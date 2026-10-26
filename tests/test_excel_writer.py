"""Tests for excel_writer module."""
import os
import tempfile
import unittest
import openpyxl

from excel_writer import _populate_reference_sheet, _base_code_expr, _modifier_expr, _dynamic_repair_formula


class TestExcelWriter(unittest.TestCase):
    def test_base_code_expr(self):
        expr = _base_code_expr(5)
        self.assertIn("$G5", expr)
        self.assertIn("NP", expr)

    def test_modifier_expr(self):
        expr = _modifier_expr(5)
        self.assertIn("$G5", expr)
        self.assertIn("SUBSTITUTE", expr)

    def test_dynamic_repair_formula(self):
        formula = _dynamic_repair_formula(10)
        self.assertIn("$G10", formula)
        self.assertIn("$H10", formula)

    def test_populate_reference_sheet(self):
        wb = openpyxl.Workbook()
        _populate_reference_sheet(wb)
        self.assertIn("Reference Data", wb.sheetnames)
        ws = wb["Reference Data"]
        # Ensure rows were populated
        self.assertGreater(ws.max_row, 1)


if __name__ == "__main__":
    unittest.main()
