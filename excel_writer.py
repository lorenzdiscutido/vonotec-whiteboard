# excel_writer.py
"""Builds and appends to the master Excel file."""
import os
import re
from collections import namedtuple

import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

from config import (
    REFERENCE_DATA, REFERENCE_SHEET_COLUMNS, MATERIAL_RULES,
    F_CRACK, F_SURFACE, F_PAINT, F_GASKET, F_SEALANT,
)
from gemini_client import MATERIALS, VALID_CODES
from image_utils import prepare_excel_image

# The Reference Data sheet range used by the lookup formulas (columns A-H: descriptions + units)
REFERENCE_RANGE = "'Reference Data'!$A:$H"


# ---------------------------------------------------------------------------
# Reference sheet
# ---------------------------------------------------------------------------

def _populate_reference_sheet(wb):
    """Creates or refreshes the Reference Data sheet with columns A-H so VLOOKUP
    formulas on Whiteboard Data can dynamically resolve units as well as causes/repairs."""
    if "Reference Data" in wb.sheetnames:
        ws_ref = wb["Reference Data"]
        ws_ref.delete_rows(1, ws_ref.max_row)
    else:
        ws_ref = wb.create_sheet(title="Reference Data")

    for row_idx, row_data in enumerate(REFERENCE_DATA, start=1):
        ws_ref.append(row_data[:REFERENCE_SHEET_COLUMNS])
        if row_idx == 1:
            for cell in ws_ref[row_idx]:
                cell.font = Font(bold=True)


# ---------------------------------------------------------------------------
# Dynamic Excel formulas watching $G{row} (Damage) and $H{row} (Dimension)
# ---------------------------------------------------------------------------

def _base_code_expr(row):
    """Excel expression that extracts the base defect code from $G{row},
    mapping 'NO PAINT' / 'NO-PAINT' to 'NP' so manual typing works either way."""
    first_word = f'LEFT(TRIM($G{row}), FIND(" ", TRIM($G{row})&" ") - 1)'
    upper_g = f'UPPER(TRIM($G{row}))'
    return f'IF(OR({upper_g}="NO PAINT", {upper_g}="NO-PAINT", {upper_g}="NOPAINT"), "NP", {first_word})'


def _modifier_expr(row):
    """Excel expression that extracts the second word (location modifier, if any) from $G{row},
    removing any hyphen (so C-F becomes CF)."""
    trimmed = f"TRIM($G{row})"
    return f'SUBSTITUTE(TRIM(MID({trimmed}, FIND(" ", {trimmed}&" ") + 1, 20)), "-", "")'


def _dynamic_repair_formula(row):
    """Builds a live Excel formula for Repair Dimension (Column O) that watches
    both $G{row} (Damage) and $H{row} (Dimension):
    - If $G{row} is cleared/empty OR $H{row} is not a number, returns "" (blank).
    - If $G{row} is edited or typed manually (e.g. C-, US, DP, FP, BP, NP, No Paint, DG, DS CC, MS FG),
      automatically applies that defect code's repair formula and repair limit."""
    h_cell = f"$H{row}"
    base = _base_code_expr(row)
    mod = _modifier_expr(row)

    crack_expr = F_CRACK.format(h=h_cell)
    surface_expr = F_SURFACE.format(h=h_cell)
    paint_expr = F_PAINT.format(h=h_cell)
    gasket_expr = F_GASKET.format(h=h_cell)
    sealant_plain = F_SEALANT.format(h=h_cell)

    # Dynamic VLOOKUP for sealant location modifier limit (falls back to plain sealant formula if no modifier)
    mod_limit = f"VLOOKUP({mod}, {REFERENCE_RANGE}, 6, FALSE)"
    sealant_with_mod = f"MROUND(MIN(({h_cell}*8/100), {mod_limit}), 0.05)"
    sealant_expr = f"IFERROR({sealant_with_mod}, {sealant_plain})"

    return (
        f'=IF(OR(TRIM($G{row})="", NOT(ISNUMBER({h_cell}))), "", '
        f'IF(OR({base}="CC-", {base}="C-", {base}="CC+", {base}="C+"), {crack_expr}, '
        f'IF(OR({base}="BH", {base}="US"), {surface_expr}, '
        f'IF(OR({base}="DP", {base}="FP", {base}="BP", {base}="NP"), {paint_expr}, '
        f'IF({base}="DG", {gasket_expr}, '
        f'IF(OR({base}="DS", {base}="MS"), {sealant_expr}, ""))))))'
    )


# ---------------------------------------------------------------------------
# Sheet layout
# ---------------------------------------------------------------------------

Layout = namedtuple("Layout", "has_unit_col unit_col photo_col photo_letter last_col title_last_col")


def _layout(has_unit_col):
    photo_col = 10 if has_unit_col else 9
    return Layout(
        has_unit_col=has_unit_col,
        unit_col=9 if has_unit_col else None,
        photo_col=photo_col,
        photo_letter=get_column_letter(photo_col),
        last_col=photo_col + 6,
        title_last_col=9 if has_unit_col else 8,
    )


def _create_master_workbook():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Whiteboard Data"

    ws.append([
        "Submitter", "Date", "Elevation", "Drop", "Floor", "Tower",
        "Defects", "", "", "Whiteboard Photo",
        "POSSIBLE CAUSE", "Possible Cause Justification", "Recommended Repair", "FINDINGS",
        "Repair Dimension", "Repair Unit",
    ])
    ws.append([
        "", "", "", "", "", "",
        "Damage", "Dimension", "Unit", "",
        "", "", "", "", "", "",
    ])

    ws.merge_cells("G1:I1")  # "Defects" spans Damage, Dimension and Unit
    ws.merge_cells("H2:I2")  # "Dimension" spans the number and its unit
    for col in ["A", "B", "C", "D", "E", "F", "J", "K", "L", "M", "N", "O", "P"]:
        ws.merge_cells(f"{col}1:{col}2")

    gray_fill = PatternFill(start_color="BFBFBF", end_color="BFBFBF", fill_type="solid")
    for row in ws["A1":"P2"]:
        for cell in row:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for ref in ("O1", "O2", "P1", "P2"):
        ws[ref].fill = gray_fill

    for letter, width in {"I": 9, "K": 24, "L": 30, "M": 28, "N": 22, "O": 18, "P": 15}.items():
        ws.column_dimensions[letter].width = width

    _populate_reference_sheet(wb)
    return wb, ws


# ---------------------------------------------------------------------------
# Splitting a dimension string ("60 CM" -> (60, "CM"))
# ---------------------------------------------------------------------------

_DIMENSION_PATTERN = re.compile(r"^\s*(\d+(?:\.\d+)?)\s*(.*)$")


def split_dimension(raw_text):
    """Splits a dimension string like '60 CM' or '120.5 CM²' into (number, unit).
    If raw_text does not start with a number, returns (raw_text, '') so nothing is lost."""
    text = str(raw_text).strip() if raw_text else ""
    if not text:
        return "", ""
    match = _DIMENSION_PATTERN.match(text)
    if not match:
        return text, ""
    num_str, unit = match.group(1), match.group(2).strip()
    number = float(num_str) if "." in num_str else int(num_str)
    return number, unit


# ---------------------------------------------------------------------------
# Rows for one whiteboard
# ---------------------------------------------------------------------------

def _expand_entries(dmg_list, dim_list):
    """Splits damage strings that hold several defect codes into one entry per code.
    Only the first code keeps the dimension."""
    expanded = []
    for i in range(max(len(dmg_list), len(dim_list))):
        dmg_str = dmg_list[i] if i < len(dmg_list) else ""
        dim_str = dim_list[i] if i < len(dim_list) else ""
        tokens = str(dmg_str).split()

        if sum(1 for t in tokens if t in VALID_CODES) > 1:
            groups, leading = [], []
            for token in tokens:
                if token in VALID_CODES:
                    groups.append(leading + [token])
                    leading = []
                elif groups:
                    groups[-1].append(token)
                else:
                    leading.append(token)
            for idx, group in enumerate(groups):
                expanded.append((" ".join(group), dim_str if idx == 0 else ""))
        else:
            expanded.append((dmg_str, dim_str))
    return expanded


def _build_rows_data(parsed_data):
    """Builds the (damage, dimension, is_title, is_data, material) tuples for one whiteboard."""
    rows_data = []
    for mat in MATERIALS:
        dmg_list = parsed_data.get(f"{mat} Damage", [])
        dim_list = parsed_data.get(f"{mat} Dimension", [])
        if not isinstance(dmg_list, list):
            dmg_list = [dmg_list] if dmg_list else []
        if not isinstance(dim_list, list):
            dim_list = [dim_list] if dim_list else []

        rows_data.append((mat, "", True, False, mat))
        expanded = _expand_entries(dmg_list, dim_list)
        if not expanded:
            rows_data.append(("", "", False, False, mat))
        else:
            for d_val, s_val in expanded:
                rows_data.append((d_val, s_val, False, True, mat))
    return rows_data


# ---------------------------------------------------------------------------
# Writing rows
# ---------------------------------------------------------------------------

def _write_title_row(ws, row, material, lay):
    cell = ws.cell(row=row, column=7)
    cell.value = material
    ws.merge_cells(start_row=row, start_column=7, end_row=row, end_column=lay.title_last_col)
    cell.font = Font(bold=True)


def _write_data_row(ws, row, g_val, h_val, material, flagged_materials, lay):
    ws.cell(row=row, column=7).value = g_val

    parts = str(g_val).split() if g_val else []
    base_code = parts[0] if parts else ""
    if str(g_val or "").upper().strip() in ("NO PAINT", "NO-PAINT", "NOPAINT"):
        base_code = "NP"
    lookup = _base_code_expr(row)

    # Column H (Dimension): plain manual-entry cell (no formula), with dynamic Unit formula in Column I
    if lay.has_unit_col:
        number, _ = split_dimension(h_val)
        ws.cell(row=row, column=8).value = number
        ws.cell(row=row, column=lay.unit_col).value = (
            f'=IF(TRIM($G{row})="", "", IFERROR(VLOOKUP({lookup}, {REFERENCE_RANGE}, 7, FALSE), ""))'
        )
    else:
        ws.cell(row=row, column=8).value = h_val

    # Lookup columns K-N (cause, justification, repair, findings) watching $G{row}
    for offset, vlookup_col in ((1, 3), (2, 4), (3, 5), (4, 2)):
        ws.cell(row=row, column=lay.photo_col + offset).value = (
            f'=IF(TRIM($G{row})="", "", IFERROR(VLOOKUP({lookup}, {REFERENCE_RANGE}, {vlookup_col}, FALSE), ""))'
        )

    # Column O (Repair Dimension) & Column P (Repair Unit): dynamically linked to $G{row} and $H{row}
    cell_repair = ws.cell(row=row, column=lay.photo_col + 5)
    cell_repair_unit = ws.cell(row=row, column=lay.photo_col + 6)
    if lay.has_unit_col:
        cell_repair.value = _dynamic_repair_formula(row)
        cell_repair.number_format = "0.00"
        cell_repair_unit.value = (
            f'=IF(TRIM($G{row})="", "", IFERROR(VLOOKUP({lookup}, {REFERENCE_RANGE}, 8, FALSE), ""))'
        )
    else:
        cell_repair.value = ""
        cell_repair_unit.value = ""

    # Red font: the material was flagged, or the code doesn't belong in this row
    is_invalid = material in flagged_materials
    if g_val and base_code not in MATERIAL_RULES.get(material, []):
        is_invalid = True
    if is_invalid:
        for col_idx in range(7, lay.last_col + 1):
            ws.cell(row=row, column=col_idx).font = Font(color="FF0000")


def write_parsed_data_to_excel(parsed_data, filepath, output_xlsx_path):
    """Adds one whiteboard's data to the master Excel file."""
    flagged_materials = set(parsed_data.get("Flagged Materials", []))

    if os.path.exists(output_xlsx_path):
        wb = openpyxl.load_workbook(output_xlsx_path)
        ws = wb["Whiteboard Data"] if "Whiteboard Data" in wb.sheetnames else wb.active
        _populate_reference_sheet(wb)
        lay = _layout(ws["J1"].value == "Whiteboard Photo")
    else:
        wb, ws = _create_master_workbook()
        lay = _layout(True)

    start_row = ws.max_row + 1
    rows_data = _build_rows_data(parsed_data)
    total_rows = len(rows_data)

    # Details (Submitter ... Tower), merged down the whole block
    for i, key in enumerate(["Submitter", "Date", "Elevation", "Drop", "Floor", "Tower"]):
        ws.cell(row=start_row, column=i + 1).value = parsed_data.get(key, "")
        if total_rows > 1:
            ws.merge_cells(start_row=start_row, start_column=i + 1,
                           end_row=start_row + total_rows - 1, end_column=i + 1)

    for i, (g_val, h_val, is_title, _is_data, material) in enumerate(rows_data):
        r = start_row + i
        if is_title:
            _write_title_row(ws, r, material, lay)
        else:
            _write_data_row(ws, r, g_val, h_val, material, flagged_materials, lay)

    # Photo, merged down the whole block
    ws.add_image(prepare_excel_image(filepath, total_rows), f"{lay.photo_letter}{start_row}")
    if total_rows > 1:
        ws.merge_cells(f"{lay.photo_letter}{start_row}:{lay.photo_letter}{start_row + total_rows - 1}")

    # Row height and alignment
    base_height = max(110 / total_rows, 20)
    for i, (_g, _h, _is_title, is_data, _mat) in enumerate(rows_data):
        r = start_row + i
        ws.row_dimensions[r].height = base_height
        for c in range(1, lay.last_col + 1):
            ws.cell(row=r, column=c).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        if lay.has_unit_col and is_data:
            ws.cell(row=r, column=lay.unit_col).alignment = Alignment(horizontal="left", vertical="center")

    ws.column_dimensions[lay.photo_letter].width = 25

    tmp_path = os.path.splitext(output_xlsx_path)[0] + "_tmp.xlsx"
    try:
        wb.save(tmp_path)
        os.replace(tmp_path, output_xlsx_path)
    except Exception:
        if os.path.exists(tmp_path):
            os.remove(tmp_path)
        raise
