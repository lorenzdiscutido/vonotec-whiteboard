# batch_processor.py
"""Processes a batch of whiteboard photos, sorts them chronologically by Date,
groups them by Day into chunks of up to 6 photos, and generates a PowerPoint (.pptx) report.

A "photo" is any object with a .name and a .getvalue() that returns the image
bytes: a Streamlit UploadedFile, or app.py's RemoteFile (which downloads the
photo from SharePoint at the moment getvalue() is called).
"""
import datetime
import hashlib
import json
import os
import re
import shutil
import threading
import traceback
import uuid
from collections import OrderedDict
from concurrent.futures import ThreadPoolExecutor, as_completed
from dataclasses import dataclass, field

from pptx import Presentation
from pptx.util import Inches

import sharepoint_client
import sync_tracker
from extraction import extract_data
from file_utils import load_processed_log, save_processed_log
from gemini_client import get_key_pool
from pptx_writer import format_defect_text, optimize_image_in_memory, write_chunks_to_pptx
from settings import ENABLE_SHAREPOINT_SYNC

MAX_PHOTOS_PER_CHUNK = 6

_DATE_FORMATS = (
    # Textual month formats first (unambiguous, e.g. "June 5, 2026", "5 June 2026")
    "%B %d, %Y",
    "%b %d, %Y",
    "%B %d %Y",
    "%b %d %Y",
    "%d %B %Y",
    "%d %b %Y",
    "%d %B, %Y",
    "%d %b, %Y",
    # Year-first ISO
    "%Y-%m-%d",
    "%Y/%m/%d",
    "%Y.%m.%d",
    # Day/Month/Year numeric (Prioritized: DD/MM/YYYY e.g. 04/07/2026 -> July 4, 2026)
    "%d/%m/%Y",
    "%d-%m-%Y",
    "%d.%m.%Y",
    "%d/%m/%y",
    "%d-%m-%y",
    # Month/Day/Year fallback
    "%m/%d/%Y",
    "%m-%d-%Y",
    "%m/%d/%y",
)


@dataclass
class BatchOutcome:
    messages: list = field(default_factory=list)  # (kind, text) tuples
    processed: int = 0
    skipped: int = 0
    created_folder: bool = False
    # One dict per photo (name, status, reason, details); status is "added", "review",
    # "skipped" or "failed". The Batch results panel draws its rows from this list.
    files: list = field(default_factory=list)
    # Messages that are not about one photo (folder created, slides generated, ...)
    notices: list = field(default_factory=list)

    def add_file(self, name, status, reason, details=""):
        self.files.append({"name": name, "status": status, "reason": reason, "details": details})

    def notice(self, kind, text):
        """A general message: kept in `messages` as before, and also listed in `notices`."""
        self.messages.append((kind, text))
        self.notices.append((kind, text))


def _parse_date(date_raw):
    """Parses the AI-extracted Date string into (sort_date, day_key, display_date).
    Valid dates sort chronologically and group by calendar day; missing/unparseable
    dates sort after valid dates."""
    raw = str(date_raw or "").strip()
    if not raw:
        return datetime.date.max, "9999-12-31", "Unknown Date"

    # Strip ordinal suffixes (1st, 2nd, 3rd, 4th), weekday names, timestamps, and normalize whitespace
    cleaned = re.sub(r"(\d+)(?:st|nd|rd|th)\b", r"\1", raw, flags=re.IGNORECASE)
    cleaned = re.sub(
        r"\b(Mon|Tue|Wed|Thu|Fri|Sat|Sun|Monday|Tuesday|Wednesday|Thursday|Friday|Saturday|Sunday)\b",
        "",
        cleaned,
        flags=re.IGNORECASE,
    )
    cleaned = re.sub(r"\b\d{1,2}:\d{2}(?::\d{2})?\b", "", cleaned)
    cleaned = re.sub(r"[|]", " ", cleaned)
    cleaned = re.sub(r"\s+", " ", cleaned).strip(" ,.-")

    for fmt in _DATE_FORMATS:
        try:
            dt = datetime.datetime.strptime(cleaned, fmt).date()
            if dt.year < 2020:
                dt = dt.replace(year=2026)
            display = f"{dt.strftime('%B')} {dt.day}, {dt.year}"
            return dt, dt.isoformat(), display
        except ValueError:
            continue

    return datetime.date.max, f"9999-12-31:{raw.lower()}", raw


def sort_and_chunk_photos(photo_records, max_per_chunk=MAX_PHOTOS_PER_CHUNK):
    """Sorts processed photo records chronologically by their AI-extracted Date,
    groups them by Date (Day), and splits each day's photos into chunks of up to
    `max_per_chunk` (default 6) photos per slide."""
    sorted_records = sorted(
        photo_records,
        key=lambda item: (item["sort_date"], item["day_key"], item["order_index"]),
    )

    grouped_by_day = OrderedDict()
    for record in sorted_records:
        day_key = record["day_key"]
        grouped_by_day.setdefault(day_key, []).append(record)

    daily_chunks = []
    for day_key, day_photos in grouped_by_day.items():
        day_label = day_photos[0]["date_display"]
        for start in range(0, len(day_photos), max_per_chunk):
            chunk_slice = day_photos[start : start + max_per_chunk]
            daily_chunks.append({
                "day_key": day_key,
                "date": day_label,
                "photos": chunk_slice,
            })

    return daily_chunks


def _create_folder(subfolder, outcome):
    """Creates the new SharePoint folder first, so it exists before uploading the presentation."""
    try:
        sharepoint_client.ensure_subfolder(subfolder)
        outcome.created_folder = True
        outcome.notice("status", f'Created the folder "{subfolder}" in SharePoint.')
    except Exception:
        traceback.print_exc()
        outcome.notice(
            "warning",
            f'Could not create the folder "{subfolder}" in SharePoint just now. '
            "The photos are still being processed and saved locally; syncing will be tried when the presentation is ready.",
        )


def _review_message(filename, parsed_data):
    """The warning shown for a photo whose automatic checks found something, or None."""
    notes = parsed_data.get("Review Notes", [])
    if not notes:
        return None
    shown = "; ".join(notes[:3])
    extra = f" (+{len(notes) - 3} more)" if len(notes) > 3 else ""
    if parsed_data.get("Flagged Materials"):
        return (
            "warning",
            f"'{filename}' was added to the PowerPoint report, but needs manual review: {shown}{extra}",
        )
    return (
        "warning",
        f"'{filename}' was added to the PowerPoint report. FYI: {shown}{extra}",
    )


def _build_file_status(name, parsed_data):
    """Builds the per-photo status dict for the Batch results UI."""
    notes = parsed_data.get("Review Notes", [])
    flagged = bool(notes) and bool(parsed_data.get("Flagged Materials"))
    extra = f" (+{len(notes) - 1} more)" if len(notes) > 1 else ""
    if flagged:
        status, reason = "review", f"Flagged for review in presentation: {notes[0]}{extra}"
    elif notes:
        status, reason = "added", f"Added to presentation. FYI: {notes[0]}{extra}"
    else:
        status, reason = "added", "Added to the PowerPoint report"
    details = "\n".join(notes)
    return {"name": name, "status": status, "reason": reason, "details": details}


def _extract_photo_record(uploaded_file, file_bytes, file_hash, order_index, outcome=None):
    """Runs AI extraction on one photo and returns its structured record, or None on error."""
    extension = os.path.splitext(uploaded_file.name)[1].lower() or ".jpg"
    temp_path = f"temp_{uuid.uuid4().hex}{extension}"
    with open(temp_path, "wb") as f:
        f.write(file_bytes)

    try:
        parsed_data = extract_data(temp_path)
        sort_date, day_key, date_display = _parse_date(parsed_data.get("Date", ""))
        drop_val = str(parsed_data.get("Drop", "") or "").strip()
        floor_val = str(parsed_data.get("Floor", "") or "").strip()
        defect_val = format_defect_text(parsed_data)

        # Pre-optimize slide image in memory for PowerPoint to eliminate end-of-batch lag
        opt_bytes = None
        opt_dims = None
        try:
            opt_stream, opt_w, opt_h = optimize_image_in_memory(file_bytes)
            opt_bytes = opt_stream.getvalue()
            opt_dims = (opt_w, opt_h)
        except Exception:
            pass

        return {
            "name": uploaded_file.name,
            "file_hash": file_hash,
            "order_index": order_index,
            "image_bytes": file_bytes,
            "optimized_image_bytes": opt_bytes,
            "optimized_dims": opt_dims,
            "parsed_data": parsed_data,
            "sort_date": sort_date,
            "day_key": day_key,
            "date": str(parsed_data.get("Date", "") or "").strip(),
            "date_display": date_display,
            "drop": drop_val,
            "floor": floor_val,
            "defect": defect_val,
        }
    except Exception as e:
        traceback.print_exc()
        if outcome is not None:
            outcome.messages.append(("error", f"An error occurred while processing {uploaded_file.name}: {e}"))
            short = str(e).strip().splitlines()[0][:90] if str(e).strip() else type(e).__name__
            outcome.add_file(uploaded_file.name, "failed", f"Extraction error: {short}", str(e))
        return None
    finally:
        if os.path.exists(temp_path):
            try:
                os.remove(temp_path)
            except Exception:
                pass


def load_existing_presentation(output_path):
    """Loads existing photo records and title block information from an existing presentation
    file (either via local .meta.json companion cache or by parsing the .pptx directly).

    Returns
    -------
    tuple[list[dict], dict]
        (existing_records, existing_title_block_info)
    """
    if not os.path.exists(output_path):
        return [], {}

    # 1. Try companion cache first if present
    meta_path = output_path + ".meta.json"
    cache_dir = output_path + ".cache"
    if os.path.exists(meta_path) and os.path.exists(cache_dir):
        try:
            with open(meta_path, "r", encoding="utf-8") as f:
                data = json.load(f)
            cached_records = data.get("records", [])
            tb_info = data.get("title_block_info", {})

            records = []
            for r in cached_records:
                img_file = r.get("image_cache_file")
                if img_file:
                    img_full_path = os.path.join(cache_dir, img_file)
                    if os.path.exists(img_full_path):
                        with open(img_full_path, "rb") as imf:
                            r["image_bytes"] = imf.read()
                sort_dt, day_key, display_date = _parse_date(r.get("date") or r.get("date_display"))
                r["sort_date"] = sort_dt
                r["day_key"] = day_key
                r["date_display"] = display_date
                if r.get("image_bytes"):
                    records.append(r)

            logo_file = tb_info.get("logo_cache_file")
            if logo_file:
                logo_path = os.path.join(cache_dir, logo_file)
                if os.path.exists(logo_path):
                    with open(logo_path, "rb") as lf:
                        tb_info["location_logo_bytes"] = lf.read()

            if records:
                return records, tb_info
        except Exception:
            pass  # Fallback to direct PPTX extraction

    # 2. Fallback: Parse the .pptx presentation directly
    try:
        prs = Presentation(output_path)
        records = []
        tb_info = {"title": "", "service_offered": "", "location_logo_bytes": None}
        tb_top_threshold = Inches(6.0)

        for slide_idx, slide in enumerate(prs.slides):
            pic_shapes = [s for s in slide.shapes if s.shape_type == 13 and s.top < tb_top_threshold]
            caption_boxes = [s for s in slide.shapes if s.has_text_frame and "Date:" in s.text and "Drop:" in s.text]

            pic_shapes.sort(key=lambda s: s.left)
            caption_boxes.sort(key=lambda s: s.left)

            for p_shape, c_box in zip(pic_shapes, caption_boxes):
                blob = p_shape.image.blob
                cap_text = c_box.text

                date_val, drop_val, floor_val, defect_val = "", "", "", ""
                for line in cap_text.splitlines():
                    line = line.strip()
                    if line.startswith("Date:"):
                        date_val = line[5:].strip()
                    elif line.startswith("Drop:"):
                        drop_val = line[5:].strip()
                    elif line.startswith("Floor:"):
                        floor_val = line[6:].strip()
                    elif line.startswith("Defect:"):
                        defect_val = line[7:].strip()

                sort_dt, day_key, display_date = _parse_date(date_val)
                p_hash = hashlib.sha256(blob).hexdigest()

                records.append({
                    "image_bytes": blob,
                    "name": f"slide_{slide_idx + 1}_{drop_val}.jpg",
                    "date_display": display_date,
                    "date": date_val,
                    "sort_date": sort_dt,
                    "day_key": day_key,
                    "drop": drop_val,
                    "floor": floor_val,
                    "defect": defect_val,
                    "parsed_data": {"Date": date_val, "Drop": drop_val, "Floor": floor_val, "Defect": defect_val},
                    "file_hash": p_hash,
                    "order_index": len(records),
                })

            if slide_idx == 0:
                for shape in slide.shapes:
                    if shape.top >= tb_top_threshold:
                        if shape.has_text_frame:
                            txt = shape.text.strip()
                            if Inches(3.2) <= shape.left < Inches(7.4):
                                tb_info["title"] = txt
                            elif Inches(7.4) <= shape.left < Inches(9.0):
                                tb_info["service_offered"] = txt
                        elif shape.shape_type == 13:
                            if Inches(8.9) <= shape.left < Inches(11.3):
                                tb_info["location_logo_bytes"] = shape.image.blob

        return records, tb_info
    except Exception:
        traceback.print_exc()
        return [], {}


def save_presentation_cache(output_path, records, title_block_info):
    """Saves a companion metadata file and cached images next to output_path
    so future batches can instantly load and merge existing photos."""
    try:
        cache_dir = output_path + ".cache"
        os.makedirs(cache_dir, exist_ok=True)
        meta_path = output_path + ".meta.json"

        cached_records = []
        for r in records:
            file_hash = r.get("file_hash") or hashlib.sha256(r.get("image_bytes", b"")).hexdigest()
            img_filename = f"{file_hash[:16]}.jpg"
            img_path = os.path.join(cache_dir, img_filename)

            if not os.path.exists(img_path) and r.get("image_bytes"):
                with open(img_path, "wb") as f:
                    f.write(r["image_bytes"])

            entry = {
                "file_hash": file_hash,
                "name": r.get("name", ""),
                "date": str(r.get("date") or ""),
                "date_display": str(r.get("date_display") or ""),
                "day_key": str(r.get("day_key") or ""),
                "drop": str(r.get("drop") or ""),
                "floor": str(r.get("floor") or ""),
                "defect": str(r.get("defect") or ""),
                "parsed_data": r.get("parsed_data") or {},
                "order_index": r.get("order_index", 0),
                "image_cache_file": img_filename,
            }
            cached_records.append(entry)

        tb_copy = {
            "title": (title_block_info or {}).get("title", ""),
            "service_offered": (title_block_info or {}).get("service_offered", ""),
        }
        logo_bytes = (title_block_info or {}).get("location_logo_bytes")
        if logo_bytes:
            logo_filename = "location_logo.png"
            with open(os.path.join(cache_dir, logo_filename), "wb") as f:
                f.write(logo_bytes)
            tb_copy["logo_cache_file"] = logo_filename

        meta_data = {
            "records": cached_records,
            "title_block_info": tb_copy,
            "total_photos": len(cached_records),
        }
        with open(meta_path, "w", encoding="utf-8") as f:
            json.dump(meta_data, f, indent=2)
    except Exception:
        pass


def run_batch(uploaded_files, output_path, subfolder="", creating_new_folder=False, on_progress=None, title_block_info=None):
    """Extracts every photo in `uploaded_files`, sorts them chronologically by Date,
    groups by Day into chunks of max 6 photos, writes the PowerPoint presentation
    to `output_path`, and syncs to SharePoint."""
    outcome = BatchOutcome()
    total = len(uploaded_files)

    def notify(i, name, finished):
        if on_progress:
            on_progress(i, total, name, finished)

    if ENABLE_SHAREPOINT_SYNC and creating_new_folder:
        _create_folder(subfolder, outcome)
    elif ENABLE_SHAREPOINT_SYNC:
        # Make sure the landing folder ("Image sorter") and the chosen subfolder exist before
        # the presentation is uploaded. Existing folders are left untouched.
        try:
            sharepoint_client.ensure_subfolder(subfolder)
        except Exception:
            traceback.print_exc()

    processed_log = load_processed_log()
    seen_hashes = set()
    valid_work_items = []
    completed_so_far = 0

    # 1. Validation & deduplication pass
    for i, uploaded_file in enumerate(uploaded_files):
        try:
            file_bytes = uploaded_file.getvalue()
        except Exception as e:
            traceback.print_exc()
            outcome.messages.append(("error", f"Could not read '{uploaded_file.name}': {e}. It was skipped."))
            outcome.add_file(uploaded_file.name, "failed", "Could not read the photo", str(e))
            completed_so_far += 1
            notify(completed_so_far - 1, uploaded_file.name, finished=True)
            continue

        file_hash = hashlib.sha256(file_bytes).hexdigest()

        if file_hash in processed_log or file_hash in seen_hashes:
            outcome.messages.append(("warning", f"Skipping '{uploaded_file.name}' - already processed."))
            outcome.add_file(
                uploaded_file.name,
                "skipped",
                "Already processed",
                "This photo's content matches one processed before (duplicates are matched by content, not by file name).",
            )
            outcome.skipped += 1
            completed_so_far += 1
            notify(completed_so_far - 1, uploaded_file.name, finished=True)
            continue

        seen_hashes.add(file_hash)
        valid_work_items.append((i, uploaded_file, file_bytes, file_hash))

    extracted_records = []
    if valid_work_items:
        key_pool = get_key_pool()
        active_slots = key_pool.active_slots()
        num_workers = max(1, min(len(active_slots), 6))

        if num_workers > 1:
            outcome_lock = threading.Lock()

            def _worker_task(u_file, f_bytes, f_hash, order_idx):
                return _extract_photo_record(u_file, f_bytes, f_hash, order_idx, outcome=None)

            with ThreadPoolExecutor(max_workers=num_workers) as executor:
                futures = {
                    executor.submit(_worker_task, u_file, f_bytes, f_hash, idx): (idx, u_file.name)
                    for (idx, u_file, f_bytes, f_hash) in valid_work_items
                }

                for future in as_completed(futures):
                    idx, fname = futures[future]
                    completed_so_far += 1
                    try:
                        rec = future.result()
                        if rec is not None:
                            with outcome_lock:
                                extracted_records.append(rec)
                        else:
                            with outcome_lock:
                                outcome.messages.append(("error", f"Extraction failed for {fname}"))
                                outcome.add_file(fname, "failed", "AI extraction returned empty", "")
                    except Exception as err:
                        traceback.print_exc()
                        with outcome_lock:
                            outcome.messages.append(("error", f"An error occurred while processing {fname}: {err}"))
                            outcome.add_file(fname, "failed", f"Worker error: {type(err).__name__}", str(err))

                    notify(completed_so_far - 1, fname, finished=True)
        else:
            for idx, u_file, f_bytes, f_hash in valid_work_items:
                notify(completed_so_far, u_file.name, finished=False)
                rec = _extract_photo_record(u_file, f_bytes, f_hash, idx, outcome=outcome)
                if rec is not None:
                    extracted_records.append(rec)
                completed_so_far += 1
                notify(completed_so_far - 1, u_file.name, finished=True)

        # Restore original input order before sorting by date
        extracted_records.sort(key=lambda r: r.get("order_index", 0))

    if not extracted_records:
        return outcome

    # Smart Append: Check if presentation already exists locally or in SharePoint
    if ENABLE_SHAREPOINT_SYNC and not os.path.exists(output_path):
        try:
            remote_filename = os.path.basename(output_path)
            sharepoint_client.download_file(remote_filename, output_path, subfolder=subfolder)
        except Exception:
            pass

    existing_records, existing_title_info = load_existing_presentation(output_path)

    if existing_records:
        merged_records = list(existing_records)
        existing_hashes = {r.get("file_hash") for r in existing_records if r.get("file_hash")}
        new_added = 0
        for r in extracted_records:
            h = r.get("file_hash")
            if not h or h not in existing_hashes:
                r_copy = dict(r)
                r_copy["order_index"] = len(existing_records) + new_added
                merged_records.append(r_copy)
                if h:
                    existing_hashes.add(h)
                new_added += 1

        all_records = merged_records

        # Inherit existing title block fields if not explicitly provided in this run
        merged_title_info = dict(existing_title_info or {})
        if title_block_info:
            if title_block_info.get("title"):
                merged_title_info["title"] = title_block_info["title"]
            if title_block_info.get("service_offered"):
                merged_title_info["service_offered"] = title_block_info["service_offered"]
            if title_block_info.get("location_logo_bytes"):
                merged_title_info["location_logo_bytes"] = title_block_info["location_logo_bytes"]
    else:
        all_records = extracted_records
        merged_title_info = title_block_info
        new_added = len(extracted_records)

    # Sort chronologically by Date, group by Day, and split into chunks of max 6 photos
    daily_chunks = sort_and_chunk_photos(all_records, max_per_chunk=MAX_PHOTOS_PER_CHUNK)
    write_chunks_to_pptx(daily_chunks, output_path, title_block_info=merged_title_info)
    save_presentation_cache(output_path, all_records, merged_title_info)

    sync_failed = False
    if ENABLE_SHAREPOINT_SYNC:
        synced, _reason = sync_tracker.sync_file(output_path, subfolder)
        if not synced:
            sync_failed = True
            outcome.notice(
                "warning",
                f"'{os.path.basename(output_path)}' was saved locally, but syncing to SharePoint failed "
                "just now. You can push it from 'Pending syncs' at the bottom of the page.",
            )

    for record in extracted_records:
        parsed_data = record["parsed_data"]
        message = _review_message(record["name"], parsed_data)
        if message:
            outcome.messages.append(message)

        file_row = _build_file_status(record["name"], parsed_data)
        if sync_failed:
            file_row["reason"] += " (SharePoint sync pending)"
            file_row["details"] = (file_row["details"] + "\n" if file_row["details"] else "") + (
                "Saved locally, but syncing to SharePoint failed just now. You can push it from 'Pending syncs'."
            )
        outcome.add_file(
            file_row["name"],
            file_row["status"],
            file_row["reason"],
            file_row["details"],
        )
        processed_log.append(record["file_hash"])
        outcome.processed += 1

    save_processed_log(processed_log)
    if existing_records:
        outcome.notice(
            "status",
            f"Appended to existing report: {len(existing_records)} existing + {new_added} new photo(s) "
            f"= {len(all_records)} total across {len(daily_chunks)} slide(s) ({len({c['day_key'] for c in daily_chunks})} day(s)).",
        )
    else:
        outcome.notice(
            "status",
            f"Generated {len(daily_chunks)} slide(s) across {len({c['day_key'] for c in daily_chunks})} day(s).",
        )
    return outcome
