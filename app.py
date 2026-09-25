import os
import json
import datetime
import base64
import streamlit as st
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill

from config import INCOMING_FOLDER, REFERENCE_DATA, MATERIAL_RULES
from image_utils import load_image, prepare_excel_image
from gemini_client import get_raw_response, parse_and_clean_json

def _populate_reference_sheet(wb):
    ws_ref = wb.create_sheet(title="Reference Data")
    header_fill = PatternFill(start_color="2F5597", end_color="2F5597", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    
    for row_idx, row_values in enumerate(REFERENCE_DATA, start=1):
        ws_ref.append(row_values)
        for col_idx in range(1, 6):
            cell = ws_ref.cell(row=row_idx, column=col_idx)
            if row_idx == 1:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            else:
                cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)

    ws_ref.column_dimensions['A'].width = 15
    ws_ref.column_dimensions['B'].width = 26
    ws_ref.column_dimensions['C'].width = 24
    ws_ref.column_dimensions['D'].width = 45
    ws_ref.column_dimensions['E'].width = 40
    ws_ref.row_dimensions[1].height = 28

def process_image_to_excel(filepath, output_xlsx_path):
    img = load_image(filepath)
    raw_text = get_raw_response(img)
    parsed_data = parse_and_clean_json(raw_text)

    if os.path.exists(output_xlsx_path):
        wb = openpyxl.load_workbook(output_xlsx_path)
        ws = wb["Whiteboard Data"] if "Whiteboard Data" in wb.sheetnames else wb.active
        if "Reference Data" not in wb.sheetnames:
            _populate_reference_sheet(wb)
    else:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Whiteboard Data"
        
        ws.append([
            "Submitter", "Date", "Elevation", "Drop", "Floor", "Tower",
            "Defects", "", "Whiteboard Photo",
            "POSSIBLE CAUSE", "Possible Cause Justification", "Recommended Repair", "FINDINGS"
        ])
        ws.append([
            "", "", "", "", "", "",
            "Damage", "Dimension", "",
            "", "", "", ""
        ])

        ws.merge_cells('G1:H1') 
        for col in ['A', 'B', 'C', 'D', 'E', 'F', 'I', 'J', 'K', 'L', 'M']:
            ws.merge_cells(f'{col}1:{col}2')
        
        gray_fill = PatternFill(start_color="BFBFBF", end_color="BFBFBF", fill_type="solid")
        for row in ws['A1':'M2']:
            for cell in row:
                cell.font = Font(bold=True)
                cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

        ws['M1'].fill = gray_fill
        ws['M2'].fill = gray_fill

        ws.column_dimensions['J'].width = 24
        ws.column_dimensions['K'].width = 30
        ws.column_dimensions['L'].width = 28
        ws.column_dimensions['M'].width = 22

        _populate_reference_sheet(wb)

    start_row = ws.max_row + 1
    VALID_CODES = {row[0] for row in REFERENCE_DATA[1:]}

    rows_data = []
    for mat in ["Sealant", "Concrete", "Paint", "Gasket"]:
        rows_data.append((mat, "", True, False, mat)) 
        
        dmg_list = parsed_data.get(f"{mat} Damage", [])
        dim_list = parsed_data.get(f"{mat} Dimension", [])
        
        expanded_dmg = []
        expanded_dim = []
        
        for i in range(max(len(dmg_list), len(dim_list))):
            dmg_str = dmg_list[i] if i < len(dmg_list) else ""
            dim_str = dim_list[i] if i < len(dim_list) else ""
            
            tokens = str(dmg_str).split()
            found_codes = [t for t in tokens if t in VALID_CODES]
            
            if len(found_codes) > 1:
                for code in found_codes:
                    expanded_dmg.append(code)
                    expanded_dim.append(dim_str) 
            else:
                expanded_dmg.append(dmg_str)
                expanded_dim.append(dim_str)
        
        max_len = max(len(expanded_dmg), 1) 
        for i in range(max_len):
            dmg_val = expanded_dmg[i] if i < len(expanded_dmg) else ""
            dim_val = expanded_dim[i] if i < len(expanded_dim) else ""
            rows_data.append((dmg_val, dim_val, False, True, mat)) 

    total_rows = len(rows_data)

    static_keys = ["Submitter", "Date", "Elevation", "Drop", "Floor", "Tower"]
    for i, key in enumerate(static_keys):
        ws.cell(row=start_row, column=i+1).value = parsed_data.get(key, "")
        if total_rows > 1:
            ws.merge_cells(start_row=start_row, start_column=i+1, end_row=start_row + total_rows - 1, end_column=i+1)

    for i, (g_val, h_val, is_title, is_data, current_mat) in enumerate(rows_data):
        r = start_row + i
        cell_g = ws.cell(row=r, column=7)
        cell_h = ws.cell(row=r, column=8)
        
        cell_g.value = g_val
        cell_h.value = h_val

        is_invalid = False
        if is_data and g_val:
            base_code = str(g_val).split()[0]
            allowed_codes = MATERIAL_RULES.get(current_mat, [])
            if base_code not in allowed_codes:
                is_invalid = True

        if is_title:
            ws.merge_cells(start_row=r, start_column=7, end_row=r, end_column=8)
            cell_g.font = Font(bold=True)
        elif is_data:
            lookup_val = f'LEFT($G{r}, FIND(" ", $G{r}&" ") - 1)'
            ws.cell(row=r, column=10).value = f'=IFERROR(VLOOKUP({lookup_val}, \'Reference Data\'!$A$2:$E$15, 3, FALSE), "")'
            ws.cell(row=r, column=11).value = f'=IFERROR(VLOOKUP({lookup_val}, \'Reference Data\'!$A$2:$E$15, 4, FALSE), "")'
            ws.cell(row=r, column=12).value = f'=IFERROR(VLOOKUP({lookup_val}, \'Reference Data\'!$A$2:$E$15, 5, FALSE), "")'
            ws.cell(row=r, column=13).value = f'=IFERROR(VLOOKUP({lookup_val}, \'Reference Data\'!$A$2:$E$15, 2, FALSE), "")'

            if is_invalid:
                for col_idx in range(7, 14):
                    ws.cell(row=r, column=col_idx).font = Font(color="FF0000")

    photo_col_letter = 'I' 
    excel_img = prepare_excel_image(filepath, total_rows)
    ws.add_image(excel_img, f"{photo_col_letter}{start_row}")
    
    if total_rows > 1:
        ws.merge_cells(f"{photo_col_letter}{start_row}:{photo_col_letter}{start_row + total_rows - 1}")
    
    base_height = max(110 / total_rows, 20)
    for i in range(total_rows):
        r = start_row + i
        ws.row_dimensions[r].height = base_height
        for c in range(1, 14):
            ws.cell(row=r, column=c).alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
            
    ws.column_dimensions[photo_col_letter].width = 25
    wb.save(output_xlsx_path)

# --- UI CONFIGURATION ---
st.set_page_config(page_title="Vonotec Whiteboard Extractor")

try:
    with open("vonotec.png", "rb") as f:
        logo_base64 = base64.b64encode(f.read()).decode()
        
    st.markdown(
        f"""
        <div style="display: flex; align-items: center; gap: 20px; margin-bottom: 10px; margin-top: -20px;">
            <div style="background-color: white; padding: 10px 15px; border-radius: 8px; display: flex; align-items: center;">
                <img src="data:image/png;base64,{logo_base64}" width="200">
            </div>
            <h1 style="margin: 0; padding: 0;">Whiteboard AI</h1>
        </div>
        """, 
        unsafe_allow_html=True
    )
except FileNotFoundError:
    st.title("Vonotec Whiteboard AI")
    st.warning("Ensure 'vonotec.png' is placed in the same folder as this script to display the logo.")

st.markdown("Upload one or multiple whiteboard photos below to automatically extract data and append it to the Master Excel File. You can also drag and drop an entire folder of images here.")

current_month_year = datetime.datetime.now().strftime("%B_%Y") 
output_xlsx_path = f"master_output_{current_month_year}.xlsx"

uploaded_files = st.file_uploader("Choose whiteboard images...", type=["jpg", "jpeg", "png"], accept_multiple_files=True)

if uploaded_files:
    st.markdown(f"**{len(uploaded_files)} image(s) selected.**")
    
    if st.button("Extract Data & Update Excel", type="primary"):
        progress_bar = st.progress(0)
        status_text = st.empty()
        
        for i, uploaded_file in enumerate(uploaded_files):
            status_text.text(f"Processing image {i + 1} of {len(uploaded_files)}: {uploaded_file.name}")
            try:
                temp_path = f"temp_{uploaded_file.name}"
                with open(temp_path, "wb") as f:
                    f.write(uploaded_file.getbuffer())
                
                process_image_to_excel(temp_path, output_xlsx_path)
                
                if os.path.exists(temp_path):
                    os.remove(temp_path)
                    
            except Exception as e:
                st.error(f"An error occurred while processing {uploaded_file.name}: {e}")
            
            progress_bar.progress((i + 1) / len(uploaded_files))
                
        status_text.text("Processing Complete.")
        st.success(f"Data from {len(uploaded_files)} file(s) successfully extracted and added to the master sheet.")

if os.path.exists(output_xlsx_path):
    st.markdown("---")
    st.subheader("Manage Master File")
    
    col1, col2 = st.columns(2)
    
    with col1:
        with open(output_xlsx_path, "rb") as file:
            st.download_button(
                label=f"Download {output_xlsx_path}",
                data=file,
                file_name=output_xlsx_path,
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                type="primary"
            )
            
    with col2:
        if st.button("Start Fresh (Clear Current Data)", type="primary"):
            os.remove(output_xlsx_path)
            st.rerun()