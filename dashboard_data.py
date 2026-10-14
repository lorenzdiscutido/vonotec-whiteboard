# dashboard_data.py
"""Reads the batch Excel files and turns them into one clean table for the
Analytics Dashboard. No Streamlit in here, so it can be tested on its own.

The sheet layout it understands (written by excel_writer.py):
    A Submitter | B Date | C Elevation | D Drop | E Floor | F Tower | G Damage | H Dimension | I Unit | J Photo ...
    ... | O Repair Dimension | P Repair Unit
    (older files have no Unit column, so their photo is in column I instead)

    Row 1-2 are headers. Every whiteboard is a block of rows: the details
    (A-F) are merged down the whole block, so only the block's first row holds
    them. Inside a block, a title row names the material (Sealant, Concrete...)
    and the rows under it are that material's defects. A row with an empty
    Damage cell is just a placeholder and is not a defect.
"""
import re
import datetime

import openpyxl
import pandas as pd

from config import REFERENCE_DATA, MATERIAL_RULES
from defect_rules import BOARD_UNITS, REPAIR_UNITS, calculate_repair

SHEET_NAME = "Whiteboard Data"
FIRST_DATA_ROW = 3
UNKNOWN = "Unknown"

COLUMNS = [
    "Folder", "File", "Submitter", "Date", "Elevation", "Drop", "Floor", "Tower",
    "Material", "Defect Code", "Defect Type", "Dimension", "Unit", "Repair Dimension", "Repair Unit",
]

# Title rows hold a material name, matched ignoring capitals
_MATERIAL_BY_LOWER = {name.lower(): name for name in MATERIAL_RULES}

# "TOWER" followed by a separator (or straight by a digit), then the tower name:
# a number (optionally with one letter) or a single letter. 'TOWERS 1' and
# 'TOWER FRONT' deliberately don't match.
_TOWER_PATTERN = re.compile(
    r"(?<![A-Za-z])TOWER(?:[\s_\-]+|(?=\d))(\d+[A-Za-z]?|[A-Za-z])(?![A-Za-z0-9])",
    re.IGNORECASE,
)

# First token of the Damage cell is the defect code, e.g. "DS CC" -> "DS"
_TYPE_BY_CODE = {row[0]: row[1] for row in REFERENCE_DATA[1:]}


# ---------------------------------------------------------------------------
# Small helpers
# ---------------------------------------------------------------------------

def _text(value):
    """Cell value as clean text ('' for empty). Whole numbers lose the '.0'."""
    if value is None:
        return ""
    if isinstance(value, datetime.datetime):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, datetime.date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _is_empty(value):
    return value is None or (isinstance(value, str) and not value.strip())


def _number(value):
    """A cell as a number, or None when it holds no number (empty, text, an Excel error)."""
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    if isinstance(value, str):
        try:
            return float(value.strip().replace(",", ""))
        except ValueError:
            return None
    return None


_DIMENSION_TEXT = re.compile(r"^\s*(\d+(?:\.\d+)?)\s*([A-Za-z]+(?:\s*(?:\^\s*2|\u00b2|2))?)?")


def _split_dimension(value):
    """A Dimension cell as (number, unit). New files hold a plain number; older ones hold
    text such as '60 CM' or '120 CM2'. (None, '') when there is no number."""
    number = _number(value)
    if number is not None:
        return number, ""
    match = _DIMENSION_TEXT.match(value) if isinstance(value, str) else None
    if not match:
        return None, ""
    unit = re.sub(r"\s*(?:\^\s*2|2)$", "\u00b2", (match.group(2) or "").upper().strip())
    return float(match.group(1)), unit


def tower_from_text(text):
    """Finds a tower in text such as a batch file name, e.g.
    'CERCA NUVEO TOWER 1- FRONT- DROP 7.xlsx' -> '1'. Returns '' when the text
    names no tower, or names more than one (it can't tell which is meant).
    Works in any capitalization, and with a space, dash or underscore after the word
    ('Tower 2', 'TOWER-2', 'tower_2', 'TOWER2'); a tower can be a number or a
    letter, optionally with a letter after the number ('3', 'A', '2B')."""
    found = {m.upper() for m in _TOWER_PATTERN.findall(str(text or ""))}
    return found.pop() if len(found) == 1 else ""


def natural_key(text):
    """Sort key so '2' comes before '10' and 'G' sits with the letters:
    'Floor 2' < 'Floor 10'."""
    return [int(part) if part.isdigit() else part.lower() for part in re.split(r"(\d+)", str(text))]


# ---------------------------------------------------------------------------
# One workbook -> rows
# ---------------------------------------------------------------------------

def parse_workbook(path, folder_label, filename):
    """Reads one batch file and returns a list of defect rows (dicts keyed by COLUMNS).
    Raises an exception if the file can't be opened or has no readable sheet."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb[SHEET_NAME] if SHEET_NAME in wb.sheetnames else wb.worksheets[0]
        ws.reset_dimensions()  # some files report a wrong size, which would hide rows

        all_rows = ws.iter_rows(min_row=1, values_only=True)
        header = next(all_rows, None)
        if header is None or len(header) < 8:
            raise ValueError("the sheet doesn't have the expected columns")
        has_unit_col = len(header) > 9 and header[9] == "Whiteboard Photo"
        repair_at = (9 if has_unit_col else 8) + 5  # index of Repair Dimension; Repair Unit is the next one

        # Workers sometimes forget to write the tower on the board, but the batch
        # name usually has it ("... TOWER 1 - FRONT - DROP 7"), so it fills the gap.
        file_tower = tower_from_text(filename)

        rows = []
        details = [""] * 6        # Submitter, Date, Elevation, Drop, Floor, Tower of the current block
        material = ""
        for index, cells in enumerate(all_rows, start=2):
            if index < FIRST_DATA_ROW:
                continue  # second header row
            cells = list(cells) + [None] * (16 - len(cells))  # pad short rows

            # A new block starts where the merged details cells hold a value
            if any(not _is_empty(c) for c in cells[:6]):
                details = [_text(c) for c in cells[:6]]
                # What is written on the board wins ("TOWER 1" is shortened to "1" so it
                # matches the others); only a blank tower is filled in from the file name.
                details[5] = tower_from_text(details[5]) or details[5] or file_tower
                material = ""

            damage = cells[6]
            if _is_empty(damage):
                continue  # placeholder row of a material with no defects, or a blank row
            damage_text = _text(damage)

            # Title row: the material's name, with nothing in the Dimension cell
            title = _MATERIAL_BY_LOWER.get(damage_text.lower())
            if title and _is_empty(cells[7]):
                material = title
                continue

            parts = damage_text.split()
            code = parts[0] if parts else ""

            # The dimension as a number (older files hold text such as "60 CM"), with its unit
            number, unit_in_text = _split_dimension(cells[7])
            unit = (_text(cells[8]) if has_unit_col else "") or unit_in_text

            # Repair dimension: the result Excel saved in the file. A file that was never opened
            # in Excel has none saved, so it is worked out here with the same formula.
            repair = _number(cells[repair_at])
            repair_unit = _text(cells[repair_at + 1])
            if repair is None:
                repair, calculated_unit = calculate_repair(parts, number)
                repair_unit = repair_unit or calculated_unit
            rows.append({
                "Folder": folder_label,
                "File": filename,
                "Submitter": details[0],
                "Date": details[1],
                "Elevation": details[2],
                "Drop": details[3],
                "Floor": details[4],
                "Tower": details[5],
                "Material": material,
                "Defect Code": code,
                "Defect Type": _TYPE_BY_CODE.get(code, ""),
                "Dimension": number,
                "Unit": unit,
                "Repair Dimension": repair,
                "Repair Unit": repair_unit,
            })
        return rows
    finally:
        wb.close()


# ---------------------------------------------------------------------------
# Rows -> table, filters and counts
# ---------------------------------------------------------------------------

def rows_to_dataframe(rows):
    """One table of every defect. Blank Tower/Floor/Material/Type become 'Unknown'
    so those defects still show up in the charts instead of vanishing."""
    df = pd.DataFrame(rows, columns=COLUMNS)
    for column in ("Floor", "Tower", "Material", "Defect Type"):
        df[column] = df[column].replace("", UNKNOWN).fillna(UNKNOWN)
    df["Dimension"] = pd.to_numeric(df["Dimension"], errors="coerce")
    df["Repair Dimension"] = pd.to_numeric(df["Repair Dimension"], errors="coerce")
    return df


def apply_filters(df, towers=None, floors=None, materials=None):
    """Keeps only the chosen values. An empty/None list for a filter means 'all'."""
    if towers:
        df = df[df["Tower"].isin(towers)]
    if floors:
        df = df[df["Floor"].isin(floors)]
    if materials:
        df = df[df["Material"].isin(materials)]
    return df


def count_by(df, column):
    """Defect frequency per value of `column`, as a table with that column and a
    'Defects' count. Floors/towers are in natural order (2 before 10); the rest
    are most-frequent first."""
    if df.empty:
        return pd.DataFrame({column: [], "Defects": []})
    counts = df.groupby(column).size().reset_index(name="Defects")
    if column in ("Floor", "Tower"):
        order = sorted(counts[column], key=natural_key)
        counts = counts.set_index(column).loc[order].reset_index()
    else:
        counts = counts.sort_values(["Defects", column], ascending=[False, True])
    return counts.reset_index(drop=True)


# Defect codes in the order of the Reference Data sheet, so tables list them the way people know them
_CODE_ORDER = {code: position for position, code in enumerate(_TYPE_BY_CODE)}

# Merge equivalent defect codes in the dashboard summary table (CC- -> C-)
_MERGED_DEFECT_CODES = {
    "CC-": "C-",
}

DEFECT_SUMMARY_COLUMNS = [
    "Defect Code", "Defect Type", "Defects", "Board Unit", "Total Dimension",
    "Repair Unit", "Total Repair Dimension", "Missing Dimension",
]


def _most_common_text(values):
    """The most frequent non-empty text in `values` ('' if there is none)."""
    counts = {}
    for value in values:
        if isinstance(value, str) and value:
            counts[value] = counts.get(value, 0) + 1
    return max(counts, key=counts.get) if counts else ""


def _convert_board_dimension(total_value, raw_unit):
    """Converts board dimensions from CM -> LM (/ 100) and CM² / SQ CM -> SQ LM (/ 10,000)."""
    unit_clean = (raw_unit or "").upper().strip()
    if total_value is None:
        if unit_clean == "CM":
            return None, "LM"
        if unit_clean in ("CM\u00b2", "CM2", "CM^2", "SQ CM", "SQ.CM", "SQ. CM"):
            return None, "SQ LM"
        return None, raw_unit

    if unit_clean == "CM":
        return float(total_value) / 100.0, "LM"
    if unit_clean in ("CM\u00b2", "CM2", "CM^2", "SQ CM", "SQ.CM", "SQ. CM"):
        return float(total_value) / 10000.0, "SQ LM"
    return float(total_value), raw_unit


def _normalize_repair_unit(raw_unit):
    """Standardizes repair units (L.M. -> LM, SQ.M. -> SQ LM) so both columns match."""
    unit_clean = (raw_unit or "").upper().strip()
    if unit_clean in ("L.M.", "L.M", "LM"):
        return "LM"
    if unit_clean in ("SQ.M.", "SQ.M", "SQM", "SQ M", "SQ LM", "SQ.L.M."):
        return "SQ LM"
    return raw_unit


def dimension_by_defect(df):
    """One row per defect code (not per material): how many defects, the total dimension written
    on the boards (converted from CM / CM² to LM / SQ LM) and the total repair dimension.
    Merges CC- into C- so both Minor Concrete Crack codes appear as one combined row."""
    if df.empty:
        return pd.DataFrame(columns=DEFECT_SUMMARY_COLUMNS)

    work_df = df.copy()
    work_df["Defect Code"] = work_df["Defect Code"].replace(_MERGED_DEFECT_CODES)

    rows = []
    for code, group in work_df.groupby("Defect Code", sort=False):
        dimensions = group["Dimension"].dropna()
        repairs = group["Repair Dimension"].dropna()

        raw_board_unit = _most_common_text(group["Unit"]) or BOARD_UNITS.get(code, "")
        raw_board_total = float(dimensions.sum()) if len(dimensions) else None
        converted_board_total, display_board_unit = _convert_board_dimension(raw_board_total, raw_board_unit)

        raw_repair_unit = _most_common_text(group["Repair Unit"]) or REPAIR_UNITS.get(code, "")
        display_repair_unit = _normalize_repair_unit(raw_repair_unit)

        rows.append({
            "Defect Code": code,
            "Defect Type": _TYPE_BY_CODE.get(code, "") or UNKNOWN,
            "Defects": len(group),
            "Board Unit": display_board_unit,
            "Total Dimension": converted_board_total,
            "Repair Unit": display_repair_unit,
            "Total Repair Dimension": float(repairs.sum()) if len(repairs) else None,
            "Missing Dimension": int(group["Dimension"].isna().sum()) if BOARD_UNITS.get(code) else 0,
        })
    summary = pd.DataFrame(rows, columns=DEFECT_SUMMARY_COLUMNS)
    summary["_order"] = summary["Defect Code"].map(lambda code: _CODE_ORDER.get(code, len(_CODE_ORDER)))
    return summary.sort_values(["_order", "Defect Code"]).drop(columns="_order").reset_index(drop=True)